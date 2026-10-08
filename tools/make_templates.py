#!/usr/bin/env python3
"""
Самые кривые описи, какие только бывают, — для проверки подбора столбцов.

Данные выдуманы. Каждый шаблон — как его делают в бухгалтериях и на объектах: формы 1С с
многоэтажной шапкой, печатные формы с нумерацией столбцов и повтором шапки на каждой
странице, таблицы без шапки, английские и киргизские заголовки, латинские буквы вместо
русских, инвентарник внутри наименования, числа текстом, таблица в углу листа, лишние листы,
формулы, ячейки без адресов, объединённые ячейки.

    python3 tools/make_templates.py      # → app/src/test/resources/templates/*.xlsx
"""
import os
import zipfile

import openpyxl

OUT = os.path.join(os.path.dirname(__file__), "..", "app", "src", "test", "resources", "templates")

ITEMS = [
    ("777/1001", "Сканер штрих-кода Honeywell", 6900, 6210),
    ("777/1002", "Сервер Dell PowerEdge R640", 450000, 300000),
    ("777/1003", "Сервер Dell PowerEdge R640", 450000, 300000),
    ("К-310", "Генератор дизельный ДГ-15", 1044321.55, 0),
    ("М 205", "Видеокамера Panasonic HC-V10", 1299017.38, 12000.5),
    ("F0000000009001", "Шкаф архивный", 15000, 15000),
    ("777/,1095", "Персональный компьютер (системный блок)", 45110, 22555),
    ("7770001007", "Ноутбук HP 250 G8", 41250, 41250),
]


def book():
    wb = openpyxl.Workbook()
    return wb, wb.active


def save(wb, name):
    wb.save(os.path.join(OUT, name))


def osv_1c():
    """ОСВ по счёту 01 из 1С: шапка документа, шапка таблицы в два этажа, группы, итоги."""
    wb, ws = book(); ws.title = "TDSheet"
    ws["A1"] = "Оборотно-сальдовая ведомость по счету 01"
    ws["A2"] = "Период: 2026 г."
    ws["A3"] = "Выводимые данные: БУ (данные бухгалтерского учета)"
    ws.append([])
    ws.append(["Основное средство", "Инв. номер", "Сальдо на начало периода", None, "Оборот за период", None, "Сальдо на конец периода", None])
    ws.append([None, None, "Дебет", "Кредит", "Дебет", "Кредит", "Дебет", "Кредит"])
    ws.merge_cells("A5:A6"); ws.merge_cells("B5:B6")
    ws.merge_cells("C5:D5"); ws.merge_cells("E5:F5"); ws.merge_cells("G5:H5")
    ws.append(["01.01"])
    ws.append(["Основное подразделение"])
    for inv, name, cost, _ in ITEMS:
        ws.append([name, inv, cost, None, None, None, cost, None])
    ws.append(["Итого", None, sum(i[2] for i in ITEMS), None, None, None, sum(i[2] for i in ITEMS), None])
    save(wb, "01-osv-1c.xlsx")


def inv1():
    """ИНВ-1: шапка в три этажа, строка «1 2 3…», повтор шапки и «Итого по странице»."""
    wb, ws = book(); ws.title = "стр.2"
    ws["A1"] = "ИНВЕНТАРИЗАЦИОННАЯ ОПИСЬ № 14 основных средств"
    ws["A3"] = "Материально ответственное лицо: Иванов И. И."
    head1 = ["Номер по порядку", "Наименование, назначение и краткая характеристика объекта", "Год выпуска (постройки, приобретения)",
             "Номер", None, None, "Фактическое наличие", None, "По данным бухгалтерского учета", None]
    head2 = [None, None, None, "инвентарный", "заводской", "паспорта", "количество", "стоимость, сом", "количество", "стоимость, сом"]
    nums = list(range(1, 11))

    def header():
        ws.append(head1); r = ws.max_row
        ws.append(head2); ws.append(nums)
        for c in "ABC":
            ws.merge_cells(f"{c}{r}:{c}{r + 1}")
        ws.merge_cells(f"D{r}:F{r}"); ws.merge_cells(f"G{r}:H{r}"); ws.merge_cells(f"I{r}:J{r}")

    ws.append([])
    header()
    for n, (inv, name, cost, _) in enumerate(ITEMS[:4], 1):
        ws.append([n, name, 2019, inv, "SN-%d" % (1000 + n), None, 1, cost, 1, cost])
    ws.append(["Итого по странице", None, None, None, None, None, 4, None, 4, None])
    header()
    for n, (inv, name, cost, _) in enumerate(ITEMS[4:], 5):
        ws.append([n, name, 2020, inv, "SN-%d" % (1000 + n), None, 1, cost, 1, cost])
    save(wb, "02-inv1.xlsx")


def no_header():
    """Ни шапки, ни заголовков — только данные."""
    wb, ws = book()
    for n, (inv, name, cost, rest) in enumerate(ITEMS, 1):
        ws.append([n, inv, name, cost, rest])
    save(wb, "03-no-header.xlsx")


def english():
    wb, ws = book(); ws.title = "Assets"
    ws.append(["No.", "Asset tag", "Description", "Location", "Qty", "Cost", "Priority"])
    for n, (inv, name, cost, _) in enumerate(ITEMS, 1):
        ws.append([n, inv, name, "Room %d" % (100 + n % 3), 1, cost, n % 3 + 1])
    save(wb, "04-english.xlsx")


def kyrgyz():
    wb, ws = book(); ws.title = "Тизме"
    ws.append(["Р/с", "Аталышы", "Инвентардык номери", "Саны", "Баасы"])
    for n, (inv, name, cost, _) in enumerate(ITEMS, 1):
        ws.append([n, name, inv, 1, cost])
    save(wb, "05-kyrgyz.xlsx")


def lookalikes():
    """Латинские буквы в заголовках, переносы строк, капс, пробелы."""
    wb, ws = book()
    ws.append(["  № п/п ", "Инв.\nнoмер", "HАИМЕНОВАНИЕ  ОС", "Стоимocть\nпервоначальная, сом", "Пpиоритет"])
    for n, (inv, name, cost, _) in enumerate(ITEMS, 1):
        ws.append([n, inv, name, cost, (n % 4) + 1])
    save(wb, "06-lookalikes.xlsx")


def inv_in_name():
    """Отдельного столбца нет: инвентарник дописан в наименование."""
    wb, ws = book()
    ws.append(["№", "Наименование", "Сумма"])
    for n, (inv, name, cost, _) in enumerate(ITEMS, 1):
        ws.append([n, f"{name}, инв. № {inv}", cost])
    save(wb, "07-inv-in-name.xlsx")


def text_numbers():
    """Деньги текстом с пробелами и «сом», номер — «1.», инвентарник числом, приоритет словами."""
    wb, ws = book()
    ws.append(["№", "Инвентарный номер", "Наименование", "Первоначальная стоимость", "Приоритет"])
    words = ["высокий", "средний", "низкий"]
    for n, (_, name, cost, _) in enumerate(ITEMS, 1):
        ws.append([f"{n}.", 130002380 + n, name, f"{cost:,.2f} сом".replace(",", " ").replace(".", ","), words[n % 3]])
    save(wb, "08-text-numbers.xlsx")


def two_blocks():
    """Два здания — два блока, у каждого свой заголовок раздела и своя шапка."""
    wb, ws = book()
    ws.append(["Здание 1 (главный корпус)"]); ws.merge_cells("A1:E1")
    ws.append(["№", "Инв. №", "Наименование", "Кол-во", "Сумма"])
    for n, (inv, name, cost, _) in enumerate(ITEMS[:4], 1):
        ws.append([n, inv, name, 1, cost])
    ws.append(["Итого по зданию 1", None, None, 4, None])
    ws.append([])
    ws.append(["Здание 2 (склад)"]); ws.merge_cells(f"A{ws.max_row}:E{ws.max_row}")
    ws.append(["№", "Инв. №", "Наименование", "Кол-во", "Сумма"])
    for n, (inv, name, cost, _) in enumerate(ITEMS[4:], 1):
        ws.append([n, inv, name, 1, cost])
    save(wb, "09-two-blocks.xlsx")


def offset():
    """Таблица в углу: с F12, пустые столбцы между, заголовок документа объединён."""
    wb, ws = book()
    ws["B2"] = "Приложение 3 к приказу № 17"; ws.merge_cells("B2:J2")
    ws["F12"] = "Наименование объекта"; ws["H12"] = "Инвентарный №"; ws["J12"] = "Балансовая стоимость"
    for n, (inv, name, cost, _) in enumerate(ITEMS, 13):
        ws[f"F{n}"] = name; ws[f"H{n}"] = inv; ws[f"J{n}"] = cost
    ws.column_dimensions["G"].hidden = True
    save(wb, "10-offset.xlsx")


def many_sheets():
    """Инструкция, сводная, пустой лист и сам список — не первым."""
    wb = openpyxl.Workbook()
    a = wb.active; a.title = "Инструкция"
    a.append(["Как заполнять опись"]); a.append(["1. Внесите наименование"]); a.append(["2. Укажите инвентарный номер"])
    s = wb.create_sheet("Сводная")
    s.append(["Подразделение", "Количество", "Сумма"]); s.append(["Главный корпус", 4, 1000]); s.append(["Склад", 4, 2000])
    wb.create_sheet("Лист1")
    l = wb.create_sheet("Список ОС")
    l.append(["№ п/п", "Инв. номер", "Наименование", "Стоимость"])
    for n, (inv, name, cost, _) in enumerate(ITEMS, 1):
        l.append([n, inv, name, cost])
    wb.save(os.path.join(OUT, "11-many-sheets.xlsx"))


def number_is_inventory():
    """«Номер» — это инвентарник, «№» — порядковый, «Код» — счёт."""
    wb, ws = book()
    ws.append(["№", "Код", "Номер", "Название", "Цена"])
    for n, (inv, name, cost, _) in enumerate(ITEMS, 1):
        ws.append([n, "01.01", inv, name, cost])
    save(wb, "12-number-is-inventory.xlsx")


def several_numbers():
    """Инвентарный, заводской, паспорт — взять инвентарный."""
    wb, ws = book()
    ws.append(["Заводской номер", "Наименование", "Номер паспорта", "Инвентарный номер", "Остаточная стоимость", "Балансовая стоимость", "Сумма износа"])
    for n, (inv, name, cost, rest) in enumerate(ITEMS, 1):
        ws.append([f"SN{9000 + n}", name, f"П-{n}", inv, rest, cost, cost - rest])
    save(wb, "13-several-numbers.xlsx")


def merged_name():
    """Наименование объединено на три столбца в каждой строке, как в печатной форме."""
    wb, ws = book()
    ws.append(["№", "Наименование", None, None, "Инв. номер", "Сумма"]); ws.merge_cells("B1:D1")
    for n, (inv, name, cost, _) in enumerate(ITEMS, 2):
        ws.append([n - 1, name, None, None, inv, cost]); ws.merge_cells(f"B{n}:D{n}")
    save(wb, "14-merged-name.xlsx")


def place_qty():
    """Место и МОЛ столбцами, количество больше одного."""
    wb, ws = book()
    ws.append(["№", "Инв. №", "Наименование", "Местонахождение", "МОЛ", "Кол-во", "Цена", "Сумма"])
    for n, (inv, name, cost, _) in enumerate(ITEMS, 1):
        ws.append([n, inv, name, f"Каб. {200 + n % 3}", "Петров П.", n % 2 + 1, cost, cost * (n % 2 + 1)])
    save(wb, "15-place-qty.xlsx")


def sections():
    """Разделы-кабинеты строками между предметами (объединены на всю ширину)."""
    wb, ws = book()
    ws.append(["№", "Инвентарный номер", "Наименование", "Стоимость"])
    for k, part in enumerate([ITEMS[:3], ITEMS[3:6], ITEMS[6:]]):
        ws.append([f"Кабинет {101 + k}"]); ws.merge_cells(f"A{ws.max_row}:D{ws.max_row}")
        for n, (inv, name, cost, _) in enumerate(part, 1):
            ws.append([n, inv, name, cost])
    save(wb, "16-sections.xlsx")


def raw_xml():
    """Как пишут сторонние программы: ячейки без адресов, общие строки с разметкой, формулы с кэшем."""
    shared = ["№", "Инв. номер", "Наименование", "Стоимость"] + [i[1] for i in ITEMS]
    def si(t):
        # Разметка внутри строки: «Сервер» жирным, остальное обычным.
        head, _, tail = t.partition(" ")
        return f'<si><r><rPr><b/></rPr><t>{head}</t></r><r><t xml:space="preserve"> {tail}</t></r></si>' if tail else f"<si><t>{t}</t></si>"
    sst = '<?xml version="1.0" encoding="UTF-8"?><sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">' + "".join(si(t) for t in shared) + "</sst>"
    rows = ['<row><c t="s"><v>0</v></c><c t="s"><v>1</v></c><c t="s"><v>2</v></c><c t="s"><v>3</v></c></row>']
    for n, (inv, name, cost, _) in enumerate(ITEMS, 1):
        rows.append(f'<row><c><f>ROW()-1</f><v>{n}</v></c><c t="inlineStr"><is><t>{inv}</t></is></c>'
                    f'<c t="s"><v>{4 + n - 1}</v></c><c><f>{cost}*1</f><v>{cost}</v></c></row>')
    sheet = '<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>' + "".join(rows) + "</sheetData></worksheet>"
    files = {
        "[Content_Types].xml": '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>',
        "_rels/.rels": '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": '<?xml version="1.0" encoding="UTF-8"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Выгрузка" sheetId="1" r:id="rId1"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="/xl/worksheets/sheet1.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>',
        "xl/worksheets/sheet1.xml": sheet,
        "xl/sharedStrings.xml": sst,
    }
    with zipfile.ZipFile(os.path.join(OUT, "17-raw-xml.xlsx"), "w", zipfile.ZIP_DEFLATED) as z:
        for k, v in files.items():
            z.writestr(k, v)


def restart_numbers():
    """Нумерация начинается заново в каждой группе, пустые строки между группами."""
    wb, ws = book()
    ws.append(["п/п", "Наименование имущества", "Инв.№", "Ед. изм.", "Кол.", "Сумма, сом"])
    for k in range(2):
        for n, (inv, name, cost, _) in enumerate(ITEMS[k * 4:(k + 1) * 4], 1):
            ws.append([n, name, inv, "шт", 1, cost])
        ws.append([])
    save(wb, "18-restart-numbers.xlsx")


def noisy():
    """Много лишних столбцов: дата ввода, счёт, амортизационная группа, примечание."""
    wb, ws = book()
    ws.append(["Счет", "Дата ввода в эксплуатацию", "Амортизационная группа", "Наименование основного средства",
               "Инвентарный номер", "Срок полезного использования", "Первоначальная стоимость", "Остаточная стоимость", "Примечание"])
    for n, (inv, name, cost, rest) in enumerate(ITEMS, 1):
        ws.append(["01.01", "15.03.2019", 3, name, inv, 60, cost, rest, "исправно" if n % 2 else ""])
    save(wb, "19-noisy.xlsx")


def no_inventory():
    """Инвентарников нет вовсе — только наименования (крайний случай: спросить)."""
    wb, ws = book()
    ws.append(["№", "Наименование", "Количество", "Цена"])
    for n, (_, name, cost, _) in enumerate(ITEMS, 1):
        ws.append([n, name, 1, cost])
    save(wb, "20-no-inventory.xlsx")


PATTERNS = ["ОС-000123", "INV-2024-0001", "0001234", "1010400123", "ВА0000123", "01.01.0023.5", "А/45-12", "КГ 123 456", "12345/2", "Н179"]
NAMES = ["Стол письменный", "Кресло офисное", "Шкаф для документов", "Принтер лазерный Canon", "Монитор Samsung 24",
         "Кондиционер настенный", "Сейф металлический", "Ноутбук Lenovo", "Проектор Epson", "Холодильник Atlant"]


def typos():
    """Заголовки с опечатками и инвентарники всех видов."""
    wb, ws = book()
    ws.append(["Н/п", "Наименавание", "Инвентраный №", "Первоночальная стоимость", "Количетсво"])
    for n, (inv, name) in enumerate(zip(PATTERNS, NAMES), 1):
        ws.append([n, name, inv, 1000 * n + 0.5, 1])
    save(wb, "21-typos-patterns.xlsx")


def patterns_no_header():
    """Без шапки, инвентарники вперемешку, перед таблицей — пустые строки."""
    wb, ws = book()
    ws.append([]); ws.append([])
    for n, (inv, name) in enumerate(zip(PATTERNS, NAMES), 1):
        ws.append([n, name, inv, 1000 * n + 0.5])
    save(wb, "22-patterns-no-header.xlsx")


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    for f in [osv_1c, inv1, no_header, english, kyrgyz, lookalikes, inv_in_name, text_numbers, two_blocks, offset,
              many_sheets, number_is_inventory, several_numbers, merged_name, place_qty, sections, raw_xml,
              restart_numbers, noisy, no_inventory, typos, patterns_no_header]:
        f()
    print(len(os.listdir(OUT)), "шаблонов в", os.path.normpath(OUT))
